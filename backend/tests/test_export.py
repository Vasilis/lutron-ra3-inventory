"""Round-trip an inventory through JSON / Markdown / CSV / XLSX exports."""

from __future__ import annotations

import io
import zipfile
from datetime import datetime, timezone

from openpyxl import load_workbook

from ra3_inventory.export import to_csv_zip, to_json, to_markdown, to_xlsx
from ra3_inventory.models import (
    Area,
    DeviceFirmwareImage,
    FirmwareInfo,
    HrefRef,
    NetworkInterface,
    ProcessorInventory,
    Project,
    RadioRa3Processor,
)
from ra3_inventory.models.button import Button, ButtonEngraving, ButtonGroup
from ra3_inventory.models.programming import Assignment, Preset, ProgrammingModel
from ra3_inventory.models.zone import Zone


def _make_inventory() -> ProcessorInventory:
    from ra3_inventory.models import SunnataHybridKeypad

    processor = RadioRa3Processor(
        href="/device/1",
        Name="Test Processor",
        DeviceType="RadioRa3Processor",
        ModelNumber="JanusProcRA3",
        SerialNumber=42,
        FirmwareImage=DeviceFirmwareImage(
            Firmware=FirmwareInfo(DisplayName="21.07.18f000"),
        ),
    )
    keypad = SunnataHybridKeypad(
        href="/device/keypad-1",
        Name="Kitchen Keypad",
        DeviceType="SunnataHybridKeypad",
        ModelNumber="RRST-HN4B-XX",
        AssociatedArea=HrefRef(href="/area/2"),
        ButtonGroups=[HrefRef(href="/buttongroup/500")],
    )
    area_root = Area(href="/area/1", Name="House")
    area_kitchen = Area(href="/area/2", Name="Kitchen", Parent=HrefRef(href="/area/1"))
    zone = Zone(
        href="/zone/100",
        Name="Kitchen Island",
        ControlType="Dimmed",
        AssociatedArea=HrefRef(href="/area/2"),
    )
    preset = Preset(
        href="/preset/200",
        Name="On 80%",
        DimmedLevelAssignments=[
            Assignment(AssignableObject=HrefRef(href="/zone/100"), Level=80),
        ],
    )
    pm = ProgrammingModel(
        href="/programmingmodel/300",
        Name="Press to 80",
        ProgrammingModelType="SimpleConditional",
        PressOnPresetAssignments=[HrefRef(href="/preset/200")],
    )
    btn = Button(
        href="/button/400",
        Name="Top Button",
        ButtonNumber=1,
        ButtonType="GeneralScene",
        Engraving=ButtonEngraving(Text="Bright"),
        ProgrammingModel=HrefRef(href="/programmingmodel/300"),
    )
    bg = ButtonGroup(
        href="/buttongroup/500",
        Name="Main",
        Buttons=[btn],
    )

    return ProcessorInventory(
        extracted_at=datetime.now(timezone.utc),
        source="fixture",
        host="192.0.2.184",
        duration_seconds=1.23,
        processor=processor,
        project=Project(
            href="/project",
            Name="Test Project",
            ProductType="Lutron RadioRA 3 Project",
        ),
        areas=[area_root, area_kitchen],
        devices=[keypad],
        zones=[zone],
        buttons=[btn],
        button_groups=[bg],
        button_group_expansions={"/device/keypad-1": [bg]},
        programming_models={"/programmingmodel/300": pm},
        presets={"/preset/200": preset},
    )


def test_json_export_round_trips() -> None:
    inv = _make_inventory()
    out = to_json(inv)
    restored = ProcessorInventory.model_validate_json(out)
    assert restored.processor.SerialNumber == 42
    assert restored.zones[0].Name == "Kitchen Island"


def test_markdown_default_is_recovery_essentials() -> None:
    """Default Markdown is the slim 'factory reset recovery' shape: no
    button-programming action chain, no hrefs in the device table."""
    inv = _make_inventory()
    md = to_markdown(inv)
    assert md.startswith("# Lutron Inventory")
    assert "## Processor" in md
    assert "## Areas" in md
    assert "## Zones" in md
    # Recovery keypad heading is engravings-only, not programming.
    assert "## Keypads — Button Engravings" in md
    assert "## Keypads — Buttons & Programming" not in md
    # Engraving labels are inlined per keypad row.
    assert "Bright" in md
    # The full PM action chain is intentionally omitted in recovery mode.
    assert "dim(Kitchen Island)=80" not in md
    # Devices table columns: Name / Type / Model / Serial / Firmware — no href.
    assert "| Name | Type | Model | Serial | Firmware |" in md


def test_markdown_verbose_emits_full_button_programming() -> None:
    """``verbose=True`` brings back the original full LEAP detail, including
    the resolved button-action chain."""
    inv = _make_inventory()
    md = to_markdown(inv, verbose=True)
    assert "## Keypads — Buttons & Programming" in md
    assert "dim(Kitchen Island)=80" in md


def test_csv_export_is_valid_zip_with_four_files() -> None:
    inv = _make_inventory()
    blob = to_csv_zip(inv)
    z = zipfile.ZipFile(io.BytesIO(blob))
    names = set(z.namelist())
    assert names == {"devices.csv", "zones.csv", "areas.csv", "buttons.csv"}
    button_csv = z.read("buttons.csv").decode("utf-8")
    assert "Bright" in button_csv  # engraving carried through
    assert "dim(Kitchen Island)=80" in button_csv  # resolved action carried through


def test_xlsx_export_is_loadable() -> None:
    inv = _make_inventory()
    blob = to_xlsx(inv)
    wb = load_workbook(io.BytesIO(blob))
    sheet_names = set(wb.sheetnames)
    assert {"Summary", "Devices", "Zones", "Areas", "Buttons"}.issubset(sheet_names)
    # Buttons sheet should have a row with our test button
    rows = list(wb["Buttons"].iter_rows(values_only=True))
    headers = rows[0]
    assert "Action" in headers
    action_col = headers.index("Action")
    actions = [r[action_col] for r in rows[1:]]
    assert any("dim(Kitchen Island)=80" in str(a) for a in actions)


def test_sanitize_inventory_redacts_identifiers_without_mutating_original() -> None:
    from ra3_inventory.sanitize import sanitize_inventory

    inv = _make_inventory()
    inv.processor.NetworkInterfaces = [NetworkInterface(MACAddress="AA:BB:CC:DD:EE:FF")]

    sanitized = sanitize_inventory(inv)

    assert sanitized.host == "192.0.2.1"
    assert sanitized.processor.SerialNumber is None
    assert sanitized.processor.Name == "Processor"
    assert sanitized.project.Name == "RA3 Inventory Demo Project"
    assert sanitized.areas[0].Name == "Area 1"
    assert inv.processor.SerialNumber == 42
    assert inv.areas[0].Name == "House"
