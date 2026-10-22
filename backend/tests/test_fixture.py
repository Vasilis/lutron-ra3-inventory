"""Validate the fully synthetic public example and its export pipeline.

The historical filename is retained for compatibility. This fixture is invented,
not a sanitized installation, and must never be replaced with a household dump.
"""

from __future__ import annotations

import csv
import io
import ipaddress
import json
import re
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from ra3_inventory.export import resolve_button_action, to_csv_zip, to_json, to_markdown, to_xlsx
from ra3_inventory.models import ProcessorInventory

FIXTURE_PATH = Path(__file__).parent.parent.parent / "examples" / "sanitized-snapshot.json"


@pytest.fixture
def synthetic_inventory() -> ProcessorInventory:
    if not FIXTURE_PATH.exists():
        pytest.skip(f"sanitized-snapshot.json not yet generated at {FIXTURE_PATH}")
    return ProcessorInventory.model_validate_json(FIXTURE_PATH.read_text())


def test_fixture_loads(synthetic_inventory: ProcessorInventory) -> None:
    """The sanitized snapshot deserializes into our schema."""
    assert synthetic_inventory.schema_version == 1
    assert synthetic_inventory.areas, "no areas in fixture?"
    assert synthetic_inventory.devices, "no devices in fixture?"
    assert synthetic_inventory.zones, "no zones in fixture?"


def test_fixture_round_trips_json(synthetic_inventory: ProcessorInventory) -> None:
    body = to_json(synthetic_inventory)
    restored = ProcessorInventory.model_validate_json(body)
    assert len(restored.devices) == len(synthetic_inventory.devices)
    assert len(restored.zones) == len(synthetic_inventory.zones)


def test_fixture_markdown_has_real_structure(synthetic_inventory: ProcessorInventory) -> None:
    md = to_markdown(synthetic_inventory)
    assert md.startswith("# Lutron Inventory")
    assert "## Processor" in md
    assert "## Summary" in md
    assert "## Areas" in md
    assert "## Zones" in md
    # Recovery-essentials default uses the engravings heading.
    assert "## Keypads — Button Engravings" in md

    # The buttons-from-expansions count should be > 0 even though the bulk
    # /button endpoint is empty on newer firmware.
    button_lines = [line for line in md.splitlines() if line.startswith("| Buttons |")]
    assert button_lines, "summary missing Buttons row"
    button_count = int(button_lines[0].split("|")[2].strip())
    assert button_count > 0, "expected non-zero button count from button_group_expansions"


def test_fixture_csv_has_buttons(synthetic_inventory: ProcessorInventory) -> None:
    blob = to_csv_zip(synthetic_inventory)
    z = zipfile.ZipFile(io.BytesIO(blob))
    button_csv = z.read("buttons.csv").decode("utf-8")
    # First non-header line should exist
    assert len(button_csv.splitlines()) > 1
    device_rows = list(csv.DictReader(io.StringIO(z.read("devices.csv").decode("utf-8"))))
    assert any(row["Firmware"] for row in device_rows)


def test_fixture_xlsx_loads(synthetic_inventory: ProcessorInventory) -> None:
    blob = to_xlsx(synthetic_inventory)
    wb = load_workbook(io.BytesIO(blob))
    assert {"Summary", "Devices", "Zones", "Areas", "Buttons"}.issubset(set(wb.sheetnames))
    rows = list(wb["Devices"].iter_rows(values_only=True))
    firmware_col = rows[0].index("Firmware")
    assert any(row[firmware_col] for row in rows[1:])


def _walk_values(value):
    yield value
    if isinstance(value, dict):
        for child in value.values():
            yield from _walk_values(child)
    elif isinstance(value, list):
        for child in value:
            yield from _walk_values(child)


def test_public_example_has_only_synthetic_metadata(synthetic_inventory) -> None:
    data = json.loads(FIXTURE_PATH.read_text())
    assert data["source"] == "fixture"
    assert ipaddress.ip_address(data["host"]) in ipaddress.ip_network("192.0.2.0/24")
    assert data["extracted_at"] == "2000-01-01T00:00:00Z"
    assert data["project"]["ProjectModifiedTimestamp"] == "2000-01-01T00:00:00Z"
    assert "Not a Real Installation" in data["project"]["Name"]
    assert data["raw_responses"] is None
    assert data["duration_seconds"] == 0
    # A small invented graph is useful for UI tests without publishing the
    # scale, field inventory or relationships of a private installation.
    assert 1 < len(data["areas"]) <= 5
    assert 1 < len(data["devices"]) <= 5
    assert 1 < len(data["zones"]) <= 5
    for value in _walk_values(data):
        if not isinstance(value, dict):
            continue
        for key, field in value.items():
            if key in {"SerialNumber", "MACAddress", "XID", "IPv4Address", "IPv4Gateway"}:
                assert field is None
            if key == "Name" and field is not None:
                assert field.startswith(("Demo ", "Synthetic "))
            if key == "Text" and field is not None:
                assert re.fullmatch(r"DEMO [A-Z]", field)
            if key == "ModelNumber" and field is not None:
                assert field.startswith("SYNTHETIC-")
            if key == "DisplayName" and field is not None:
                assert field == "synthetic-1.0"
            if key == "Installed":
                assert field is None
    assert synthetic_inventory.processor.NetworkInterfaces == []


def test_public_example_graph_references_resolve(synthetic_inventory) -> None:
    data = json.loads(FIXTURE_PATH.read_text())
    resources = {}
    references = set()
    for value in _walk_values(data):
        if not isinstance(value, dict) or "href" not in value:
            continue
        href = value["href"]
        references.add(href)
        if len(value) > 1:
            assert href not in resources, f"duplicate resource definition: {href}"
            resources[href] = value
        if href != "/project":
            assert re.fullmatch(r"/[a-z]+/9[0-9]{5}", href), href
    assert references <= resources.keys(), references - resources.keys()
    devices = {d.href: d for d in synthetic_inventory.devices}
    zones = {z.href: z for z in synthetic_inventory.zones}
    for zone in zones.values():
        device = devices[zone.Device.href]
        assert zone.href in {ref.href for ref in device.LocalZones}
        assert zone.AssociatedArea == device.AssociatedArea
    areas = {area.href: area for area in synthetic_inventory.areas}
    for area in areas.values():
        seen = set()
        current = area
        while current.Parent is not None:
            assert current.href not in seen, "area cycle"
            seen.add(current.href)
            current = areas[current.Parent.href]
    for device_href, groups in synthetic_inventory.button_group_expansions.items():
        device = devices[device_href]
        assert {g.href for g in groups} == {ref.href for ref in device.ButtonGroups}
        for group in groups:
            assert group.Parent.href == device_href
            for button in group.Buttons:
                assert button.Parent.href == group.href
                assert button.ProgrammingModel.href in synthetic_inventory.programming_models


def test_published_markdown_is_current_synthetic_render(synthetic_inventory) -> None:
    published = FIXTURE_PATH.with_name("sanitized-inventory.md").read_text()
    assert published == to_markdown(synthetic_inventory)
    assert "Synthetic Inventory Example" in published
    # Exercise both the modern and legacy programming shapes, including
    # resolved zone targets, rather than leaving the demonstration graph empty.
    zones = {zone.href: zone for zone in synthetic_inventory.zones}
    actions = [
        resolve_button_action(
            button, synthetic_inventory.programming_models, synthetic_inventory.presets, zones
        )
        for groups in synthetic_inventory.button_group_expansions.values()
        for group in groups
        for button in group.Buttons
    ]
    assert any("press:dim(Demo Load A)=50" in action for action in actions)
    assert any("press:switch(Demo Load B)=On" in action for action in actions)
