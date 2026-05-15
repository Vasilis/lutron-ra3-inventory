"""Excel (.xlsx) export — multi-sheet workbook via openpyxl."""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..models import ProcessorInventory, area_path
from ._resolver import resolve_button_action

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="2C3E50")
_WRAP = Alignment(wrap_text=True, vertical="top")


def _write_sheet(ws, headers: list[str], rows: list[list]) -> None:
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = _HEADER_FONT
        c.fill = _HEADER_FILL
    for ridx, row in enumerate(rows, start=2):
        for cidx, val in enumerate(row, start=1):
            cell = ws.cell(row=ridx, column=cidx, value=val)
            cell.alignment = _WRAP

    # Auto-size columns (rough — based on header length and longest row)
    for cidx in range(1, len(headers) + 1):
        width = max(
            len(str(headers[cidx - 1])),
            *(len(str(row[cidx - 1])) for row in rows if len(row) >= cidx),
            default=10,
        )
        ws.column_dimensions[get_column_letter(cidx)].width = min(width + 2, 60)
    ws.freeze_panes = "A2"


def to_xlsx(inv: ProcessorInventory) -> bytes:
    """Return a multi-sheet .xlsx workbook for one ProcessorInventory."""
    wb = Workbook()
    wb.remove(wb.active)

    areas_by_href = {a.href: a for a in inv.areas}
    zones_by_href = {z.href: z for z in inv.zones}
    devices_by_href = {d.href: d for d in inv.devices}

    # Summary
    summary = wb.create_sheet("Summary")
    _write_sheet(
        summary,
        ["Field", "Value"],
        [
            ["Project", inv.project.Name or ""],
            ["ProductType", inv.project.ProductType or ""],
            ["Project Modified", inv.project.ProjectModifiedTimestamp or ""],
            ["Extracted At", inv.extracted_at.isoformat()],
            ["Host", inv.host],
            ["Schema", f"v{inv.schema_version}"],
            ["Partial", str(inv.partial)],
            ["", ""],
            ["Areas", len(inv.areas)],
            ["Devices", len(inv.devices)],
            ["Zones", len(inv.zones)],
            ["Buttons", len(inv.buttons)],
            ["LEDs", len(inv.leds)],
            ["Virtual Buttons", len(inv.virtual_buttons)],
            ["Area Scenes", len(inv.area_scenes)],
            ["Programming Models", len(inv.programming_models)],
            ["Presets", len(inv.presets)],
            ["Timeclock Rules", len(inv.timeclock_event_rules)],
        ],
    )

    # Devices
    devices = wb.create_sheet("Devices")
    rows = []
    for d in inv.devices:
        fw_name = ""
        if d.FirmwareImage and d.FirmwareImage.Firmware:
            fw_name = d.FirmwareImage.Firmware.DisplayName or ""
        ahref = d.AssociatedArea.href if d.AssociatedArea else None
        rows.append([
            d.href, d.Name or "", d.DeviceType, d.ModelNumber or "",
            str(d.SerialNumber or ""), fw_name, d.AddressedState or "",
            area_path(ahref, areas_by_href) if ahref else "",
        ])
    _write_sheet(
        devices,
        ["href", "Name", "DeviceType", "ModelNumber", "SerialNumber",
         "Firmware", "AddressedState", "Area"],
        rows,
    )

    # Zones
    zones = wb.create_sheet("Zones")
    rows = []
    for z in inv.zones:
        cat = f"{z.Category.Type or ''}/{z.Category.SubType or ''}" if z.Category else ""
        rows.append([
            z.href, z.Name or "", z.ControlType or "", cat,
            z.Device.href if z.Device else "",
        ])
    _write_sheet(zones, ["href", "Name", "ControlType", "Category", "Device"], rows)

    # Areas
    areas = wb.create_sheet("Areas")
    rows = []
    for a in inv.areas:
        rows.append([
            a.href, a.Name or "",
            a.Parent.href if a.Parent else "",
            area_path(a.href, areas_by_href),
        ])
    _write_sheet(areas, ["href", "Name", "Parent", "FullPath"], rows)

    # Buttons
    buttons = wb.create_sheet("Buttons")
    rows = []
    for dhref, bgs in inv.button_group_expansions.items():
        d = devices_by_href.get(dhref)
        for bg in bgs:
            for btn in bg.Buttons or []:
                action = resolve_button_action(
                    btn, inv.programming_models, inv.presets, zones_by_href
                )
                eng = btn.Engraving.Text if btn.Engraving else ""
                rows.append([
                    dhref, d.Name if d else "", btn.href,
                    btn.ButtonNumber or "", eng or "", btn.Name or "",
                    btn.ButtonType or "", action,
                ])
    _write_sheet(
        buttons,
        ["device_href", "device_name", "button_href", "ButtonNumber",
         "Engraving", "Name", "ButtonType", "Action"],
        rows,
    )

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
